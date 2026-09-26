import streamlit as st
import os
import re
import smtplib
import time
import imaplib
import json
import hashlib
import string
import email
from email.header import decode_header
from email.utils import parseaddr
from filelock import FileLock
from email.mime.text import MIMEText
from email.mime.multipart import MIMEMultipart
from email.mime.application import MIMEApplication
import openpyxl

try:
    import pandas as pd
except ImportError:
    pd = None

# ==========================================
# 網頁基本設定
# ==========================================
st.set_page_config(page_title="公關寄信系統", layout="wide")

USERS_FILE = "users.json"
USERS_LOCK = "users.json.lock"
PROJECTS_FILE = "projects.json"
PROJECTS_LOCK = "projects.json.lock"

def decode_str(s):
    if not s: return ""
    try:
        value, charset = decode_header(s)[0]
        if charset:
            return value.decode(charset)
        elif isinstance(value, bytes):
            return value.decode('utf-8', errors='ignore')
        return str(value)
    except Exception:
        return str(s)

class SafeFormatter(string.Formatter):
    def get_value(self, key, args, kwargs):
        if isinstance(key, str):
            return kwargs.get(key, f"{{{key}}}")
        return super().get_value(key, args, kwargs)

def safe_format_template(template, context_dict):
    if not template: return ""
    try:
        formatter = SafeFormatter()
        return formatter.format(template, **context_dict)
    except Exception:
        res = str(template)
        for k, v in context_dict.items():
            res = res.replace(f"{{{k}}}", str(v))
        return res

def hash_password(password):
    return hashlib.sha256(password.encode()).hexdigest()

def safe_replace_file(src, dst, max_retries=5):
    for attempt in range(max_retries):
        try:
            os.replace(src, dst)
            return
        except PermissionError:
            if attempt == max_retries - 1: raise
            time.sleep(0.02)

def load_users():
    with FileLock(USERS_LOCK, timeout=10):
        if not os.path.exists(USERS_FILE):
            default_users = {"admin": {"password": hash_password("admin123"), "role": "admin", "real_name": "系統管理員"}}
            temp_file = f"{USERS_FILE}.tmp"
            with open(temp_file, "w", encoding="utf-8") as f:
                json.dump(default_users, f, ensure_ascii=False, indent=4)
            safe_replace_file(temp_file, USERS_FILE)
            return default_users
        with open(USERS_FILE, "r", encoding="utf-8") as f:
            try: return json.load(f)
            except Exception: return {}

def save_users(users_data):
    with FileLock(USERS_LOCK, timeout=10):
        temp_file = f"{USERS_FILE}.tmp"
        with open(temp_file, "w", encoding="utf-8") as f:
            json.dump(users_data, f, ensure_ascii=False, indent=4)
        safe_replace_file(temp_file, USERS_FILE)

def load_projects():
    with FileLock(PROJECTS_LOCK, timeout=10):
        if not os.path.exists(PROJECTS_FILE): return {}
        with open(PROJECTS_FILE, "r", encoding="utf-8") as f:
            try: 
                data = json.load(f)
                # 舊資料相容性升級
                for p_name, p_data in data.items():
                    if "mode" not in p_data: p_data["mode"] = "individual"
                    if "replies" not in p_data: p_data["replies"] = []
                return data
            except Exception: return {}

def save_projects(projects_data):
    with FileLock(PROJECTS_LOCK, timeout=10):
        temp_file = f"{PROJECTS_FILE}.tmp"
        with open(temp_file, "w", encoding="utf-8") as f:
            json.dump(projects_data, f, ensure_ascii=False, indent=4)
        safe_replace_file(temp_file, PROJECTS_FILE)

def extract_email(contact_str):
    if not contact_str or str(contact_str).lower() in ("nan", "none", "null"): return None
    match = re.search(r'[a-zA-Z0-9_.+-]+@[a-zA-Z0-9-]+\.[a-zA-Z0-9-.]+', str(contact_str))
    return match.group(0) if match else None

def read_excel_data(uploaded_file):
    if pd is not None:
        try:
            df = pd.read_excel(uploaded_file, sheet_name="全部彙總清單")
            return df.to_dict(orient="records")
        except Exception:
            try:
                df = pd.read_excel(uploaded_file, sheet_name=0)
                return df.to_dict(orient="records")
            except Exception: pass
    try:
        wb = openpyxl.load_workbook(uploaded_file, data_only=True)
        sheet_name = "全部彙總清單" if "全部彙總清單" in wb.sheetnames else wb.sheetnames[0]
        ws = wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True))
        if not rows: return []
        headers = [str(h).strip() if h is not None else "" for h in rows[0]]
        data = []
        for row in rows[1:]:
            if any(c is not None for c in row):
                data.append({headers[i]: (row[i] if i < len(row) and row[i] is not None else "") for i in range(len(headers))})
        return data
    except Exception as e:
        st.error(f"❌ Excel 讀取失敗：{e}")
        return []

# 所有預設字串留空
DEFAULT_TEMPLATE = ""

if "logged_in" not in st.session_state:
    st.session_state.logged_in = False
    st.session_state.username = ""
    st.session_state.role = ""
    st.session_state.real_name = ""
if "page" not in st.session_state: st.session_state.page = "home"
if "current_project" not in st.session_state: st.session_state.current_project = None
if "gmail_account" not in st.session_state: st.session_state.gmail_account = ""
if "gmail_password" not in st.session_state: st.session_state.gmail_password = ""

if not st.session_state.logged_in:
    st.title("🔐 公關寄信系統")
    st.markdown("請輸入您的專屬帳號與密碼以登入系統。")
    with st.form("login_form"):
        login_user = st.text_input("帳號").strip()
        login_pwd = st.text_input("密碼", type="password").strip()
        if st.form_submit_button("登入", type="primary"):
            users_db = load_users()
            if login_user in users_db and users_db[login_user]["password"] == hash_password(login_pwd):
                st.session_state.logged_in = True
                st.session_state.username, st.session_state.role = login_user, users_db[login_user]["role"]
                st.session_state.real_name = users_db[login_user].get("real_name", login_user)
                st.rerun()
            else:
                st.error("⚠️ 帳號或密碼錯誤！")
    st.stop()

projects_db = load_projects()

unread_count = 0
for p_data in projects_db.values():
    for reply in p_data.get("replies", []):
        if not reply.get("read", True) and (st.session_state.role == "admin" or reply.get("owner") == st.session_state.real_name):
            unread_count += 1

st.sidebar.markdown(f"👤 登入者：**{st.session_state.real_name}**")
if st.sidebar.button("🚪 登出", use_container_width=True):
    st.session_state.logged_in = False
    st.rerun()
st.sidebar.divider()

nav_options = ["🏠 專案與寄信區", f"📥 收件與回信匣 {'🔴' if unread_count > 0 else ''}"]
if st.session_state.role == "admin": nav_options.append("⚙️ 系統後台管理")
app_mode = st.sidebar.radio("📌 系統功能導覽", nav_options)
st.sidebar.divider()

if "⚙️ 系統後台管理" in app_mode:
    st.title("⚙️ 系統後台管理")
    tab1, tab2 = st.tabs(["👥 帳號管理與重設密碼", "📊 團隊寄件總覽"])
    with tab1:
        col_new, col_reset = st.columns(2)
        users_db = load_users()
        with col_new:
            st.subheader("➕ 建立新帳號")
            with st.form("add_user_form"):
                n_usr = st.text_input("登入帳號").strip()
                n_name = st.text_input("成員姓名").strip()
                n_pwd = st.text_input("預設密碼", type="password").strip()
                n_role = st.selectbox("帳號權限", ["user (一般)", "admin (管理員)"])
                if st.form_submit_button("新增帳號", type="primary", use_container_width=True):
                    if n_usr in users_db: st.error("帳號已存在！")
                    elif not n_usr or not n_pwd: st.error("帳號密碼不得為空！")
                    else:
                        users_db[n_usr] = {"password": hash_password(n_pwd), "role": "admin" if "admin" in n_role else "user", "real_name": n_name or n_usr}
                        save_users(users_db)
                        st.success(f"成功建立帳號：{n_name}")
                        st.rerun()
        with col_reset:
            st.subheader("🔑 重設成員密碼")
            with st.form("reset_pwd_form"):
                target_user = st.selectbox("選擇要重設密碼的帳號", list(users_db.keys()))
                new_pwd = st.text_input("輸入新密碼", type="password").strip()
                if st.form_submit_button("強制重設密碼", type="primary", use_container_width=True):
                    if not new_pwd: st.error("新密碼不得為空！")
                    else:
                        users_db[target_user]["password"] = hash_password(new_pwd)
                        save_users(users_db)
                        st.success(f"✅ 已重設 {users_db[target_user].get('real_name', target_user)} 的密碼！")
                        st.rerun()
        st.table([{"登入帳號": u, "真實姓名": d.get("real_name", u), "權限": d.get("role", "user")} for u, d in users_db.items()])
    with tab2:
        st.subheader("📂 專案寄件紀錄總覽 (上帝視角)")
        for p_name, p_data in projects_db.items():
            sent_list = p_data.get("sent_companies", [])
            with st.expander(f"📁 {p_name} (團隊共寄出 {len(sent_list)} 封)"):
                for record in sent_list:
                    if isinstance(record, dict): st.markdown(f"- **{record.get('company', '?')}** (Email: {record.get('email', '未記錄')} | 寄件負責人: **{record.get('sender', '?')}**)")
                    else: st.markdown(f"- **{record}** (早期紀錄)")

elif "📥 收件與回信匣" in app_mode:
    st.title("📥 廠商回信與通知中心")
    st.subheader("1. 郵件伺服器認證")
    col1, col2, col3 = st.columns([3, 3, 2])
    test_email = col1.text_input("團隊 Gmail 信箱", value=st.session_state.gmail_account).strip()
    test_pwd = col2.text_input("應用程式密碼", value=st.session_state.gmail_password, type="password").strip()
    reply_folder = col3.text_input("自動歸檔資料夾", value="FRC_Replies")
    
    if st.button("🔄 強制掃描近期回信", type="primary", use_container_width=True):
        if not test_email or not test_pwd: st.error("請輸入信箱與密碼！")
        else:
            with st.spinner("🚀 正在掃描 Gmail 近 3 天內信件..."):
                try:
                    mail = imaplib.IMAP4_SSL("imap.gmail.com", timeout=20)
                    mail.login(test_email, test_pwd)
                    status, _ = mail.select(reply_folder)
                    if status != 'OK': mail.create(reply_folder)
                    mail.select("INBOX")
                    
                    date_limit = (pd.Timestamp.now() - pd.Timedelta(days=3)).strftime("%d-%b-%Y")
                    status, messages = mail.search(None, f'(SINCE "{date_limit}")')
                    
                    if status == "OK" and messages[0]:
                        msg_nums = messages[0].split()
                        sent_map = {}
                        for p_name, p_data in projects_db.items():
                            for record in p_data.get("sent_companies", []):
                                if isinstance(record, dict) and record.get("email"):
                                    clean_email = str(record["email"]).strip().lower()
                                    sent_map[clean_email] = (p_name, record.get("company"), record.get("sender"))

                        new_reply_count = 0
                        if len(msg_nums) > 100: msg_nums = msg_nums[-100:]

                        for num in msg_nums:
                            res, header_data = mail.fetch(num, "(BODY.PEEK[HEADER.FIELDS (FROM SUBJECT)])")
                            for response_part in header_data:
                                if isinstance(response_part, tuple):
                                    header_msg = email.message_from_bytes(response_part[1])
                                    from_header = decode_str(header_msg.get("From"))
                                    _, addr = parseaddr(from_header)
                                    addr_lower = str(addr).strip().lower()
                                    subject_check = decode_str(header_msg.get("Subject"))
                                    
                                    if addr_lower in sent_map:
                                        proj_name, company_name, original_sender = sent_map[addr_lower]
                                        is_duplicate = any(r["subject"] == subject_check and r["email"] == addr_lower for r in projects_db[proj_name].get("replies", []))
                                                
                                        if not is_duplicate:
                                            res, full_msg_data = mail.fetch(num, "(RFC822)")
                                            for full_response_part in full_msg_data:
                                                if isinstance(full_response_part, tuple):
                                                    msg = email.message_from_bytes(full_response_part[1])
                                                    body = "無法解析文字內容"
                                                    if msg.is_multipart():
                                                        for part in msg.walk():
                                                            if part.get_content_type() == "text/plain":
                                                                try: body = part.get_payload(decode=True).decode('utf-8', errors='ignore'); break
                                                                except: pass
                                                    else:
                                                        try: body = msg.get_payload(decode=True).decode('utf-8', errors='ignore')
                                                        except: pass
                                                    
                                                    clipped_body = body[:1000] + ("\n...(內容過長已省略)" if len(body)>1000 else "")
                                                    projects_db[proj_name]["replies"].append({
                                                        "company": company_name, "email": addr, "subject": subject_check,
                                                        "body": clipped_body, "read": False, "time": time.strftime("%Y-%m-%d %H:%M"), "owner": original_sender 
                                                    })
                                                    new_reply_count += 1
                                                    mail.copy(num, reply_folder)
                                                    mail.store(num, '+FLAGS', '\\Deleted')
                        mail.expunge()
                        if new_reply_count > 0:
                            save_projects(projects_db)
                            st.success(f"🎉 成功攔截 {new_reply_count} 封新回信！")
                            time.sleep(1.5)
                            st.rerun()
                        else: st.info("沒有發現未處理過的回信。")
                    else: st.info("近 3 天內沒有收到任何信件。")
                    mail.logout()
                except Exception as e: st.error(f"連線失敗：{e}")

    st.divider()
    st.subheader("2. 廠商回信匣")
    has_any_reply = False
    for p_name, p_data in projects_db.items():
        replies = p_data.get("replies", [])
        visible_replies = [(idx, reply) for idx, reply in enumerate(replies) if st.session_state.role == "admin" or reply.get("owner") == st.session_state.real_name]
        if not visible_replies: continue
            
        has_any_reply = True
        st.markdown(f"#### 📂 專案：{p_name}")
        for real_idx, reply in reversed(visible_replies):
            is_unread = not reply.get("read", True)
            icon = "🔴" if is_unread else "🟢"
            with st.expander(f"{icon} 來自 {reply['company']} ({reply['email']}) - {reply['subject']}"):
                st.caption(f"接收時間：{reply.get('time', '未知')} | 負責人：{reply.get('owner', '未知')}")
                st.text(reply['body'])
                if is_unread:
                    if st.button("標示為已讀", key=f"read_{p_name}_{real_idx}"):
                        projects_db[p_name]["replies"][real_idx]["read"] = True
                        save_projects(projects_db)
                        st.rerun()
    if not has_any_reply: st.info("目前尚無屬於您的廠商回信紀錄。")

elif "🏠 專案與寄信區" in app_mode:
    if st.session_state.page == "home":
        st.title("✉️ 公關寄信系統 - 專案大廳")
        col_new, col_list = st.columns([1, 2])
        
        with col_new:
            st.subheader("➕ 建立新專案")
            with st.form("new_project_form"):
                new_proj_name = st.text_input("專案命名：", placeholder="填入專案名稱...").strip()
                proj_mode = st.radio("寄件模式", ["個別寄送 (每家廠商個別發送)", "一鍵群發 (使用統一模板寄給全部廠商)"])
                if st.form_submit_button("建立專案", type="primary", use_container_width=True):
                    if not new_proj_name: st.error("名稱不能為空！")
                    elif new_proj_name in projects_db: st.error("專案已存在！")
                    else:
                        mode_val = "bulk" if "群發" in proj_mode else "individual"
                        projects_db[new_proj_name] = {"sent_companies": [], "template": DEFAULT_TEMPLATE, "replies": [], "mode": mode_val}
                        save_projects(projects_db)
                        st.rerun()

        with col_list:
            st.subheader("📂 現有專案列表")
            for proj_name, proj_data in list(projects_db.items()):
                sent_companies = proj_data.get('sent_companies', [])
                display_count = len(sent_companies) if st.session_state.role == "admin" else sum(1 for r in sent_companies if isinstance(r, dict) and r.get('sender') == st.session_state.real_name)
                with st.container(border=True):
                    c1, c2, c3 = st.columns([6, 2, 2])
                    with c1:
                        mode_label = "【群發模式】" if proj_data.get("mode") == "bulk" else "【個別模式】"
                        st.markdown(f"#### {proj_name} {mode_label}")
                        st.caption(f"{'團隊共寄出' if st.session_state.role == 'admin' else '您已寄出'} {display_count} 封")
                    with c2:
                        if st.button("📂 開啟", key=f"open_{proj_name}", use_container_width=True):
                            st.session_state.current_project = proj_name
                            st.session_state.page = "project"
                            st.rerun()
                    with c3:
                        if st.session_state.role == "admin" and st.button("🗑️ 刪除", key=f"del_{proj_name}", use_container_width=True):
                            del projects_db[proj_name]
                            save_projects(projects_db)
                            st.rerun()

    elif st.session_state.page == "project":
        current_proj = st.session_state.current_project
        proj_data = projects_db[current_proj]

        if st.sidebar.button("🔙 返回專案大廳", type="primary", use_container_width=True):
            st.session_state.page = "home"
            st.rerun()
            
        st.sidebar.divider()
        st.sidebar.header("🔐 寄件帳號設定")
        sender_email = st.sidebar.text_input("團隊 Gmail 信箱", value=st.session_state.gmail_account).strip()
        sender_password = st.sidebar.text_input("應用程式密碼", value=st.session_state.gmail_password, type="password").strip()
        st.session_state.gmail_account, st.session_state.gmail_password = sender_email, sender_password
        
        st.sidebar.header("📂 系統設定")
        pdf_dir = st.sidebar.text_input("本機附件資料夾名稱", value="企劃書檔案").strip()
        backup_folder = st.sidebar.text_input("Gmail 寄件備份標籤", value="FRC_Sponsorship").strip()

        st.title(f"📁 專案：{current_proj} {'(🚀 一鍵群發模式)' if proj_data.get('mode') == 'bulk' else ''}")
        st.divider()

        st.header("Step 1: 團隊與信件格式")
        col1, col2, col3 = st.columns(3)
        team_name = col1.text_input("團隊名稱", value="", placeholder="例如：FRC 團隊名稱").strip()
        contact_person = col2.text_input("聯絡人", value="", placeholder="例如：公關長 王小明").strip()
        contact_phone = col3.text_input("電話", value="", placeholder="例如：0912-345-678").strip()

        email_template = st.text_area("✏️ 信件內容", value=proj_data.get("template", ""), height=200, placeholder="請輸入統一的信件內容，可使用大括號插入變數，例如 {企業／贊助單位}")
        if email_template != proj_data.get("template"):
            projects_db[current_proj]["template"] = email_template
            save_projects(projects_db)

        st.header("Step 2 & 3: 載入與寄出")
        uploaded_file = st.file_uploader("上傳名單 (.xlsx)", type=["xlsx"])
        
        sent_list = proj_data.get("sent_companies", [])
        all_sent_companies = [r.get("company") if isinstance(r, dict) else r for r in sent_list]
        my_sent_companies = [r.get("company") for r in sent_list if isinstance(r, dict) and (st.session_state.role == "admin" or r.get("sender") == st.session_state.real_name)]
        
        if uploaded_file is not None:
            records = read_excel_data(uploaded_file)
            if records:
                # 判斷是否為群發模式
                if proj_data.get("mode") == "bulk":
                    st.subheader("🚀 一鍵群發寄信作業")
                    unsent_records = [r for r in records if r.get('企業／贊助單位') and r.get('企業／贊助單位') not in all_sent_companies]
                    
                    if not unsent_records:
                        st.success("🎉 Excel 表單中的所有廠商皆已寄出！")
                    else:
                        st.info(f"尚有 {len(unsent_records)} 家廠商未寄送。點擊下方按鈕將一次寄出給所有未寄送的廠商。")
                        
                        # 預覽第一筆
                        st.markdown("### 📝 群發信件預覽 (以第一筆為例)")
                        first_record = unsent_records[0]
                        first_company = first_record.get('企業／贊助單位')
                        first_id = str(first_record.get('編號', '000')).zfill(3)
                        format_dict = dict(first_record)
                        format_dict.update({"team_name": team_name, "contact_person": contact_person, "contact_phone": contact_phone, "pdf_filename": f"{first_id}_{first_company}_贊助企劃書.pdf"})
                        st.text(safe_format_template(email_template, format_dict))
                        
                        if st.button(f"🚀 一鍵寄出給剩下 {len(unsent_records)} 家廠商", type="primary"):
                            if not sender_email or not sender_password:
                                st.error("⚠️ 帳號未設定！")
                            else:
                                progress_bar = st.progress(0)
                                status_text = st.empty()
                                
                                try:
                                    server = smtplib.SMTP("smtp.gmail.com", 587, timeout=15)
                                    server.starttls()
                                    server.login(sender_email, sender_password)
                                    
                                    try:
                                        imap = imaplib.IMAP4_SSL("imap.gmail.com", timeout=15)
                                        imap.login(sender_email, sender_password)
                                        if imap.select(backup_folder)[0] != 'OK': imap.create(backup_folder)
                                    except:
                                        imap = None
                                        
                                    success_count = 0
                                    for idx, row_data in enumerate(unsent_records):
                                        company = row_data.get('企業／贊助單位')
                                        to_email = extract_email(row_data.get('聯絡資訊', ''))
                                        status_text.text(f"正在寄送 ({idx+1}/{len(unsent_records)}): {company} ...")
                                        
                                        if to_email:
                                            pdf_filename = f"{str(row_data.get('編號', '000')).zfill(3)}_{company}_贊助企劃書.pdf"
                                            pdf_path = os.path.join(pdf_dir, pdf_filename)
                                            
                                            fmt_dict = dict(row_data)
                                            fmt_dict.update({"team_name": team_name, "contact_person": contact_person, "contact_phone": contact_phone, "pdf_filename": pdf_filename})
                                            preview_text = safe_format_template(email_template, fmt_dict)
                                            
                                            msg = MIMEMultipart()
                                            msg['From'], msg['To'], msg['Subject'] = sender_email, to_email, f"【贊助合作邀請】{team_name} — 敬致 {company}"
                                            msg.attach(MIMEText(preview_text, 'plain', 'utf-8'))
                                            if os.path.exists(pdf_path):
                                                with open(pdf_path, 'rb') as f:
                                                    attach = MIMEApplication(f.read(), _subtype="pdf")
                                                    attach.add_header('Content-Disposition', 'attachment', filename=pdf_filename)
                                                    msg.attach(attach)
                                            
                                            server.send_message(msg)
                                            if imap:
                                                try: imap.append(backup_folder, '\\Seen', imaplib.Time2Internaldate(time.time()), msg.as_bytes())
                                                except: pass
                                                
                                            projects_db[current_proj]["sent_companies"].append({
                                                "company": company, "sender": st.session_state.real_name, "email": to_email
                                            })
                                            success_count += 1
                                            time.sleep(1) # 暫停1秒避免被 Gmail 判定為垃圾信
                                            
                                        progress_bar.progress((idx + 1) / len(unsent_records))
                                        
                                    server.quit()
                                    if imap: imap.logout()
                                    save_projects(projects_db)
                                    status_text.text(f"✅ 群發完成！成功寄出 {success_count} 封。")
                                    time.sleep(2)
                                    st.rerun()
                                except Exception as e:
                                    st.error(f"❌ 群發中斷：{e}")

                else:
                    # 個別寄送模式 (原本的邏輯)
                    company_list = list(set([r.get('企業／贊助單位', '') for r in records if r.get('企業／贊助單位')]))
                    selected_company = st.selectbox("🔍 選擇廠商", company_list) if company_list else None
                    
                    if selected_company:
                        row_data = next((r for r in records if r.get('企業／贊助單位') == selected_company), {})
                        to_email = extract_email(row_data.get('聯絡資訊', ''))
                        
                        if to_email: st.success(f"📧 **即將寄出至 (目標信箱)：** `{to_email}`")
                        else: st.error("⚠️ **警告：** 在 Excel 中找不到此廠商的有效 Email，將無法寄送！")
                        
                        pdf_filename = f"{str(row_data.get('編號', '000')).zfill(3)}_{selected_company}_贊助企劃書.pdf"
                        pdf_path = os.path.join(pdf_dir, pdf_filename)
                        
                        format_dict = dict(row_data)
                        format_dict.update({"team_name": team_name, "contact_person": contact_person, "contact_phone": contact_phone, "pdf_filename": pdf_filename})
                        
                        preview_text = safe_format_template(email_template, format_dict)
                        st.text(preview_text)
                        
                        is_sent = selected_company in all_sent_companies
                        
                        if st.button("🚀 再次寄送" if is_sent else "🚀 確定寄送", type="primary"):
                            if not sender_email or not sender_password or not to_email:
                                st.error("⚠️ 帳號未設定，或廠商無有效 Email！")
                            else:
                                with st.spinner('寄送信件中...'):
                                    try:
                                        server = smtplib.SMTP("smtp.gmail.com", 587, timeout=10)
                                        server.starttls()
                                        server.login(sender_email, sender_password)
                                        
                                        msg = MIMEMultipart()
                                        msg['From'], msg['To'], msg['Subject'] = sender_email, to_email, f"【贊助合作邀請】{team_name} — 敬致 {selected_company}"
                                        msg.attach(MIMEText(preview_text, 'plain', 'utf-8'))
                                        
                                        if os.path.exists(pdf_path):
                                            with open(pdf_path, 'rb') as f:
                                                attach = MIMEApplication(f.read(), _subtype="pdf")
                                                attach.add_header('Content-Disposition', 'attachment', filename=pdf_filename)
                                                msg.attach(attach)
                                        
                                        server.send_message(msg)
                                        server.quit()
                                        
                                        try:
                                            imap = imaplib.IMAP4_SSL("imap.gmail.com", timeout=10)
                                            imap.login(sender_email, sender_password)
                                            if imap.select(backup_folder)[0] != 'OK': imap.create(backup_folder)
                                            imap.append(backup_folder, '\\Seen', imaplib.Time2Internaldate(time.time()), msg.as_bytes())
                                            imap.logout()
                                        except: pass
                                        
                                        if not is_sent:
                                            proj_data["sent_companies"].append({
                                                "company": selected_company,
                                                "sender": st.session_state.real_name,
                                                "email": to_email
                                            })
                                            save_projects(projects_db)
                                        
                                        st.success("✅ 寄出成功！")
                                        time.sleep(1.5)
                                        st.rerun()
                                    except Exception as e:
                                        st.error(f"❌ 寄件失敗：{e}")